"""Report generation: PDF reconciliation reports, Excel mismatch workbook, audit evidence pack."""
import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    AuditLog, ExceptionHistory, MigrationRun, ReconException, ReconSummary, RecordResult, Report, SignOff,
)
from . import ai, analytics

PRIMARY = colors.HexColor("#0070AD")
SECONDARY = colors.HexColor("#004C7F")

REPORT_TYPES = {
    "BUSINESS_RECONCILIATION": ("pdf", "Business Reconciliation Report"),
    "MIGRATION_SUMMARY": ("pdf", "Migration Summary"),
    "COMPLIANCE": ("pdf", "Compliance Report"),
    "MANAGEMENT_SUMMARY": ("pdf", "Management Summary"),
    "DETAILED_MISMATCH": ("xlsx", "Detailed Mismatch Report"),
    "AUDIT_PACK": ("zip", "Audit Pack"),
}
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "zip": "application/zip",
}
COMPLIANCE_DOMAINS = ["aml", "kyc", "sanctions", "customer"]


def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("H1x", parent=styles["Heading1"], textColor=SECONDARY, fontSize=18))
    styles.add(ParagraphStyle("H2x", parent=styles["Heading2"], textColor=PRIMARY, fontSize=13))
    styles.add(ParagraphStyle("Cell", parent=styles["BodyText"], fontSize=7.5, leading=9))
    return styles


def _table(rows: list[list], col_widths=None) -> Table:
    styles = _styles()
    data = [[Paragraph(str(c), styles["Cell"]) for c in row] for row in rows]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8D3DD")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7FB")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 8 * mm, "FICrime Migration & Reconciliation Platform - CONFIDENTIAL")
    canvas.drawRightString(doc.pagesize[0] - 15 * mm, 8 * mm, f"Page {doc.page}")
    canvas.restoreState()


def _exceptions(db: Session, run_id: int, domains: list[str] | None = None) -> list[ReconException]:
    stmt = select(ReconException).where(ReconException.run_id == run_id)
    if domains:
        stmt = stmt.where(ReconException.domain.in_(domains))
    return list(db.scalars(stmt.order_by(ReconException.code)))


def build_pdf(db: Session, run: MigrationRun, report_type: str, user: str) -> bytes:
    styles = _styles()
    title = REPORT_TYPES[report_type][1]
    domains = COMPLIANCE_DOMAINS if report_type == "COMPLIANCE" else None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    story = [Paragraph(title, styles["H1x"]),
             Paragraph(f"{run.name} &nbsp;|&nbsp; {run.source_system} &rarr; {run.target_system} &nbsp;|&nbsp; "
                       f"{run.environment} &nbsp;|&nbsp; Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} by {user}",
                       styles["BodyText"]), Spacer(1, 6)]

    excs = _exceptions(db, run.id, domains)
    open_excs = [e for e in excs if e.status in {"DETECTED", "ASSIGNED", "INVESTIGATING"}]
    kpis = [["Records processed", "Reconciliation %", "Exceptions", "Open", "Critical open", "Sign-off"],
            [f"{run.records_processed:,}", f"{run.reconciliation_pct:.2f}%", len(excs), len(open_excs),
             sum(e.severity == "CRITICAL" for e in open_excs), run.signoff_status]]
    story += [_table(kpis), Spacer(1, 8), Paragraph("Executive summary", styles["H2x"]),
              Paragraph(ai.executive_narrative(db, run), styles["BodyText"]), Spacer(1, 6)]

    stats = [s for s in analytics.domain_stats(db, [run.id]) if not domains or s["domain"] in domains]
    story += [Paragraph("Reconciliation by domain (Level 1 / Level 2)", styles["H2x"]), _table(
        [["Domain", "Source", "Target", "Matched", "Mismatched", "Missing in target", "Missing in source",
          "Match %", "Exceptions", "Open"]] +
        [[s["domain"], s["source_count"], s["target_count"], s["MATCHED"], s["MISMATCHED"],
          s["MISSING_IN_TARGET"], s["MISSING_IN_SOURCE"],
          "-" if s["match_pct"] is None else f"{s['match_pct']}%", s["exceptions"], s["open_exceptions"]]
         for s in stats])]

    if report_type in {"BUSINESS_RECONCILIATION", "COMPLIANCE", "MIGRATION_SUMMARY"}:
        summ = list(db.scalars(select(ReconSummary).where(ReconSummary.run_id == run.id)
                               .order_by(ReconSummary.domain, ReconSummary.id)))
        summ = [s for s in summ if not domains or s.domain in domains]
        story += [Spacer(1, 6), Paragraph("Level 1 control totals", styles["H2x"]), _table(
            [["Domain", "Metric", "Source", "Target", "Difference", "Result"]] +
            [[s.domain, s.metric, f"{s.source_value:,.2f}", f"{s.target_value:,.2f}", f"{s.difference:+,.2f}",
              "MATCH" if s.matched else "BREAK"] for s in summ])]

    if report_type != "MIGRATION_SUMMARY":
        rules = [r for r in analytics.rule_results(db, [run.id]) if not domains or r.domain in domains]
        story += [Spacer(1, 6), Paragraph("Business rule results", styles["H2x"]), _table(
            [["Rule", "Category", "Domain", "Actual", "Expected", "Result"]] +
            [[r.rule_name, r.category, r.domain, r.actual, r.expected, "PASS" if r.passed else "FAIL"]
             for r in rules])]

    story += [Spacer(1, 6), Paragraph("Exceptions by category and root cause", styles["H2x"])]
    pivot: dict[tuple, int] = {}
    for e in excs:
        pivot[(e.category, e.root_cause, e.severity)] = pivot.get((e.category, e.root_cause, e.severity), 0) + 1
    story.append(_table([["Category", "Root cause", "Severity", "Count"]] +
                        [[*k, v] for k, v in sorted(pivot.items(), key=lambda kv: -kv[1])]))

    if report_type in {"BUSINESS_RECONCILIATION", "COMPLIANCE"}:
        crit = [e for e in excs if e.severity in {"CRITICAL", "HIGH"}][:60]
        story += [PageBreak(), Paragraph("Critical & high exceptions (first 60)", styles["H2x"]), _table(
            [["ID", "Domain", "Severity", "Root cause", "Description", "Owner", "Status"]] +
            [[e.code, e.domain, e.severity, e.root_cause, e.description, e.owner or "-", e.status] for e in crit],
            col_widths=[30 * mm, 20 * mm, 18 * mm, 35 * mm, 110 * mm, 20 * mm, 22 * mm])]

    signoffs = list(db.scalars(select(SignOff).where(SignOff.run_id == run.id).order_by(SignOff.at)))
    story += [Spacer(1, 6), Paragraph("Business sign-off", styles["H2x"]), _table(
        [["Area", "Decision", "By", "Role", "Comment", "At"]] +
        ([[s.area, s.decision, s.actor, s.role, s.comment, f"{s.at:%Y-%m-%d %H:%M}"] for s in signoffs]
         or [["-", "No sign-offs recorded", "", "", "", ""]]))]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def build_excel(db: Session, run: MigrationRun) -> bytes:
    wb = Workbook()
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0070AD")

    def sheet(name: str, headers: list[str], rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append([("" if v is None else v) for v in row])
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, idx).column_letter].width = max(12, min(45, len(h) + 6))

    wb.remove(wb.active)
    sheet("Summary", ["Domain", "Source", "Target", "Matched", "Mismatched", "Missing in target",
                      "Missing in source", "Match %", "Exceptions"],
          [[s["domain"], s["source_count"], s["target_count"], s["MATCHED"], s["MISMATCHED"],
            s["MISSING_IN_TARGET"], s["MISSING_IN_SOURCE"], s["match_pct"], s["exceptions"]]
           for s in analytics.domain_stats(db, [run.id])])
    sheet("L1 Totals", ["Domain", "Metric", "Source", "Target", "Difference", "Matched"],
          [[s.domain, s.metric, s.source_value, s.target_value, s.difference, s.matched]
           for s in db.scalars(select(ReconSummary).where(ReconSummary.run_id == run.id))])
    records = list(db.scalars(select(RecordResult).where(RecordResult.run_id == run.id,
                                                         RecordResult.status != "MATCHED")
                              .order_by(RecordResult.domain, RecordResult.record_key)))
    sheet("L2 Record Breaks", ["Domain", "Record key", "Status", "Fields different", "Total abs variance"],
          [[r.domain, r.record_key, r.status, len(r.field_differences or []), r.variance] for r in records])
    sheet("L3 Field Differences", ["Domain", "Record key", "Field", "Source value", "Target value", "Difference"],
          [[r.domain, r.record_key, d["field"], str(d["source"]) if d["source"] is not None else None,
            str(d["target"]) if d["target"] is not None else None, d["difference"]]
           for r in records for d in (r.field_differences or [])])
    sheet("Exceptions", ["Exception ID", "Domain", "Category", "Severity", "Root cause", "Record key", "Field",
                         "Source", "Target", "Variance", "Owner", "Status", "Created", "Description"],
          [[e.code, e.domain, e.category, e.severity, e.root_cause, e.record_key, e.field, e.source_value,
            e.target_value, e.variance, e.owner, e.status, e.created_at.strftime("%Y-%m-%d %H:%M"), e.description]
           for e in _exceptions(db, run.id)])
    sheet("Rule Results", ["Rule", "Category", "Domain", "Passed", "Actual", "Expected", "Detail"],
          [[r.rule_name, r.category, r.domain, r.passed, r.actual, r.expected, r.detail]
           for r in analytics.rule_results(db, [run.id])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(headers: list[str], rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode()


def build_audit_pack(db: Session, run: MigrationRun, user: str) -> bytes:
    excs = _exceptions(db, run.id)
    exc_ids = [e.id for e in excs]
    history = list(db.scalars(select(ExceptionHistory).where(ExceptionHistory.exception_id.in_(exc_ids))
                              .order_by(ExceptionHistory.at))) if exc_ids else []
    code_by_id = {e.id: e.code for e in excs}
    audit = [a for a in db.scalars(select(AuditLog).order_by(AuditLog.at))
             if (a.entity == "migration_run" and a.entity_id == str(run.id))
             or (a.entity == "exception" and a.entity_id.startswith(f"EXC-{run.id:04d}-"))
             or (a.entity == "report" and a.detail.get("run_id") == run.id)]
    files = {
        "01_reconciliation_report.pdf": build_pdf(db, run, "BUSINESS_RECONCILIATION", user),
        "02_detailed_mismatch_report.xlsx": build_excel(db, run),
        "03_exception_log.csv": _csv(
            ["exception_id", "domain", "category", "severity", "root_cause", "record_key", "field",
             "source_value", "target_value", "variance", "owner", "status", "resolution", "created_at"],
            [[e.code, e.domain, e.category, e.severity, e.root_cause, e.record_key, e.field, e.source_value,
              e.target_value, e.variance, e.owner, e.status, e.resolution, e.created_at.isoformat()] for e in excs]),
        "04_exception_workflow_history.csv": _csv(
            ["exception_id", "from_status", "to_status", "actor", "comment", "at"],
            [[code_by_id[h.exception_id], h.from_status, h.to_status, h.actor, h.comment, h.at.isoformat()]
             for h in history]),
        "05_approval_log.csv": _csv(
            ["area", "decision", "actor", "role", "comment", "at"],
            [[s.area, s.decision, s.actor, s.role, s.comment, s.at.isoformat()]
             for s in db.scalars(select(SignOff).where(SignOff.run_id == run.id).order_by(SignOff.at))]),
        "06_audit_log.csv": _csv(["actor", "action", "entity", "entity_id", "detail", "at"],
                                 [[a.actor, a.action, a.entity, a.entity_id, json.dumps(a.detail), a.at.isoformat()]
                                  for a in audit]),
        "07_migration_evidence.json": json.dumps({
            "run": {"id": run.id, "name": run.name, "source_system": run.source_system,
                    "target_system": run.target_system, "environment": run.environment, "status": run.status,
                    "records_processed": run.records_processed, "reconciliation_pct": run.reconciliation_pct,
                    "signoff_status": run.signoff_status,
                    "started_at": run.started_at.isoformat() if run.started_at else None,
                    "completed_at": run.completed_at.isoformat() if run.completed_at else None},
            "domains": analytics.domain_stats(db, [run.id]),
            "control_totals": analytics.summary_metrics(db, [run.id]),
            "rules": [{"rule": r.rule_name, "domain": r.domain, "passed": r.passed, "actual": r.actual,
                       "expected": r.expected} for r in analytics.rule_results(db, [run.id])],
        }, indent=2, default=str).encode(),
    }
    manifest = {"generated_at": datetime.now(timezone.utc).isoformat(), "generated_by": user, "run_id": run.id,
                "files": {name: {"sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}
                          for name, data in files.items()}}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
        zf.writestr("manifest.json", json.dumps(manifest, indent=2))
    return buf.getvalue()


def generate(db: Session, run: MigrationRun, report_type: str, user: str) -> Report:
    fmt, title = REPORT_TYPES[report_type]
    if fmt == "pdf":
        content = build_pdf(db, run, report_type, user)
    elif fmt == "xlsx":
        content = build_excel(db, run)
    else:
        content = build_audit_pack(db, run, user)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    report = Report(run_id=run.id, report_type=report_type, file_format=fmt,
                    filename=f"{report_type.lower()}_run{run.id}_{stamp}.{fmt}", content_type=CONTENT_TYPES[fmt],
                    sha256=hashlib.sha256(content).hexdigest(), size_bytes=len(content), content=content,
                    created_by=user)
    db.add(report)
    db.flush()
    return report
